"""
Locked report builder for episodic retention analysis xlsx deliverables.

Import build_report() and call it once per show, after you've run
diagnose.py and hand-written the spike_rows / overall_bullets / headline
based on judgment (this script does NOT decide what's a spike or what
matters — see SKILL.md for that).

    from report_template import build_report
    build_report(
        show_name="...",
        csv_path="...",
        out_path="....xlsx",
        headline="...",           # one bold sentence, the punchline
        spike_rows=[(ep, event, drop_str, note), ...],  # 3-5 rows, self-contained
        overall_bullets=[...],    # supporting narrative
        chart_markers=[...],      # episode numbers to pointer-mark on every chart
        show_id_filter=None,      # pass if the CSV is mixed with another show's rows
        tab_path=DEFAULT_TAB_PATH,        # override if your benchmark file lives elsewhere
        tab_show_id=DEFAULT_TAB_SHOW_ID,
    )
"""
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.marker import Marker
from openpyxl.chart.label import DataLabelList
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

PURPLE = "4A3AA7"
BLUE = "2A78D6"
ORANGE = "EB6834"
GREEN = "0E8A5C"
RED = "C0392B"
LIGHT_GREY = "FAF9F7"
RULE = "D9D8D2"

DEFAULT_TAB_PATH = "/Users/tejas/Downloads/TAB Data/TAB_FULL_701eps.csv"
DEFAULT_TAB_SHOW_ID = "cbc35fa13bbe154873ec97f78f2b4cf4b66c7d31"


def fnum(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def load(path, show_id_filter=None):
    with open(path, encoding="utf-8-sig") as f:
        rows = list(csv.DictReader(f))
    data = []
    for r in rows:
        if show_id_filter and r["Show ID"] != show_id_filter:
            continue
        drop = r.get("CB D0+ Drop% overall")
        if drop in (None, ""):
            drop = r.get("CB D0+ Drop% d7")
        data.append({
            "ep": int(r["Story No."]),
            "title": r["Story Title"],
            "ldau": int(r["Total LDAU"]),
            "drop": fnum(drop),
            "ma10": fnum(r["Moving Avg 10"]),
            "ma50": fnum(r["Moving Avg 50"]),
            "launch": r["Launch Date"],
        })
    data.sort(key=lambda r: r["ep"])
    return data


def build_report(show_name, csv_path, out_path, headline, spike_rows, overall_bullets,
                  chart_markers, subtitle_extra="", footnote_extra="", show_id_filter=None,
                  tab_path=DEFAULT_TAB_PATH, tab_show_id=DEFAULT_TAB_SHOW_ID):
    """
    spike_rows: list of (episode_label, event, drop_str, note).
        note must be self-contained — fold any "fix this" call directly into
        the text, never reference a separate recommendations section (there
        isn't one, by design — see SKILL.md).
    chart_markers: episode numbers (ints) to highlight with red pointer
        markers + episode-number data labels on every chart. Pick the worst
        point of each actionable cluster; skip clusters marked "no action
        needed" in spike_rows.
    show_id_filter: pass when csv_path's export is mixed with TAB (or
        another show's) rows — always check with diagnose.py first.
    """
    show = load(csv_path, show_id_filter)
    tab = load(tab_path, tab_show_id)
    baseline = show[50]["ldau"]
    tab_baseline = tab[50]["ldau"]
    tab_by_ep = {d["ep"]: d for d in tab}
    final_ep = show[-1]["ep"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Analysis"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "portrait"
    ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.5, bottom=0.5)

    widths = [12, 22, 16, 56]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    thin = Side(style="thin", color=RULE)
    border_all = Border(left=thin, right=thin, top=thin, bottom=thin)

    r = 1
    ws.merge_cells(f"A{r}:D{r}")
    c = ws.cell(row=r, column=1, value=f"{show_name} — Episodic Retention Analysis")
    c.font = Font(size=15, bold=True, color=PURPLE)
    r += 1
    ws.merge_cells(f"A{r}:D{r}")
    c = ws.cell(row=r, column=1,
                value=f"Overall cohort · {final_ep} episodes · Benchmarked vs The Alpha's Bride (701 eps) · {subtitle_extra}")
    c.font = Font(size=10, italic=True, color="444444")
    r += 1
    ws.merge_cells(f"A{r}:D{r}")
    c = ws.cell(row=r, column=1, value=headline)
    c.font = Font(size=11, bold=True, color="1a1a1a")
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 30
    r += 2

    # ---- KEY SPIKES (top of page, always — no separate recommendations section) ----
    ws.merge_cells(f"A{r}:D{r}")
    c = ws.cell(row=r, column=1, value="KEY SPIKES")
    c.font = Font(size=11, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=PURPLE)
    c.alignment = Alignment(indent=1)
    ws.row_dimensions[r].height = 20
    r += 1

    headers = ["Episode", "Event", "Drop (×100) / vs TAB", "What's happening"]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=r, column=i, value=h)
        c.font = Font(bold=True, size=10, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=PURPLE)
        c.border = border_all
    r += 1

    for i, (ep, event, drop, note) in enumerate(spike_rows):
        fill = PatternFill("solid", fgColor=LIGHT_GREY) if (i % 2 == 1) else PatternFill("solid", fgColor="FFFFFF")
        c1 = ws.cell(row=r, column=1, value=ep); c1.font = Font(bold=True, size=10); c1.border = border_all; c1.fill = fill
        c1.alignment = Alignment(wrap_text=True, vertical="top")
        c2 = ws.cell(row=r, column=2, value=event); c2.font = Font(size=10); c2.border = border_all; c2.fill = fill
        c2.alignment = Alignment(wrap_text=True, vertical="top")
        c3 = ws.cell(row=r, column=3, value=drop); c3.font = Font(bold=True, size=9.5, color=RED); c3.border = border_all; c3.fill = fill
        c3.alignment = Alignment(wrap_text=True, vertical="top")
        c4 = ws.cell(row=r, column=4, value=note); c4.font = Font(size=9.5); c4.border = border_all; c4.fill = fill
        c4.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 66
        r += 1

    ws.merge_cells(f"A{r}:D{r}")
    c = ws.cell(row=r, column=1,
                value="Ep 1 excluded (launch sampling always expected). Final episode shows a blank drop% — "
                      "end-of-data artifact, kept as LDAU but excluded from the drop series. Benchmark ratio = "
                      "show block-drop ÷ TAB block-drop at the same episode range, both baselined at Ep 51; "
                      ">1x = losing audience faster than TAB. " + footnote_extra)
    c.font = Font(size=9, italic=True, color="767671")
    c.alignment = Alignment(wrap_text=True)
    ws.row_dimensions[r].height = 26
    r += 2

    # ---- OVERALL (supporting narrative only) ----
    ws.merge_cells(f"A{r}:D{r}")
    c = ws.cell(row=r, column=1, value="OVERALL")
    c.font = Font(size=11, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=BLUE)
    c.alignment = Alignment(indent=1)
    ws.row_dimensions[r].height = 20
    r += 1

    for b in overall_bullets:
        ws.merge_cells(f"A{r}:D{r}")
        c = ws.cell(row=r, column=1, value="•  " + b)
        c.font = Font(size=10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 32
        r += 1

    r += 1
    chart1_row = r; r += 18
    chart2_row = r; r += 18
    chart3_row = r; r += 18

    # ================= Sheet 2: Raw Data =================
    ws2 = wb.create_sheet("Raw Data")
    ws2.sheet_view.showGridLines = False
    rd_headers = ["Episode", "Story Title", "LDAU", "Drop %", "Drop x100", "MA10", "MA10 x100", "MA50", "MA50 x100",
                  "TAB LDAU (ep-aligned)", "Retention Index (Ep51=100)", "TAB Retention Index (Ep51=100)",
                  "Spike Drop x100", "Spike LDAU", "Spike Retention Index"]
    for i, h in enumerate(rd_headers, start=1):
        c = ws2.cell(row=1, column=i, value=h)
        c.font = Font(bold=True, size=10, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=PURPLE)
        c.border = border_all
    widths2 = [10, 34, 12, 11, 11, 11, 11, 11, 11, 16, 18, 20, 14, 12, 18]
    for i, w in enumerate(widths2, start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.freeze_panes = "A2"

    marker_set = set(chart_markers)

    for idx, d in enumerate(show):
        row = idx + 2
        ws2.cell(row=row, column=1, value=d["ep"])
        ws2.cell(row=row, column=2, value=d["title"])
        ws2.cell(row=row, column=3, value=d["ldau"])
        drop, ma10, ma50 = d["drop"], d["ma10"], d["ma50"]
        c4 = ws2.cell(row=row, column=4, value=drop)
        if drop is not None: c4.number_format = "0.00%"
        drop_x100 = drop * 100 if drop is not None else None
        c5 = ws2.cell(row=row, column=5, value=drop_x100)
        if drop is not None: c5.number_format = "0.00"
        c6 = ws2.cell(row=row, column=6, value=ma10)
        if ma10 is not None: c6.number_format = "0.00%"
        c7 = ws2.cell(row=row, column=7, value=(ma10 * 100 if ma10 is not None else None))
        if ma10 is not None: c7.number_format = "0.00"
        c8 = ws2.cell(row=row, column=8, value=ma50)
        if ma50 is not None: c8.number_format = "0.00%"
        c9 = ws2.cell(row=row, column=9, value=(ma50 * 100 if ma50 is not None else None))
        if ma50 is not None: c9.number_format = "0.00"
        tab_row = tab_by_ep.get(d["ep"])
        tab_ldau = tab_row["ldau"] if tab_row else None
        ws2.cell(row=row, column=10, value=tab_ldau)
        ret_idx = None
        if d["ep"] >= 51:
            ret_idx = round(100 * d["ldau"] / baseline, 2)
            ws2.cell(row=row, column=11, value=ret_idx)
            if tab_ldau:
                ws2.cell(row=row, column=12, value=round(100 * tab_ldau / tab_baseline, 2))
        is_marker = d["ep"] in marker_set
        ws2.cell(row=row, column=13, value=(drop_x100 if (is_marker and drop_x100 is not None) else None))
        ws2.cell(row=row, column=14, value=(d["ldau"] if is_marker else None))
        ws2.cell(row=row, column=15, value=(ret_idx if (is_marker and ret_idx is not None) else None))
        for col in range(1, 16):
            ws2.cell(row=row, column=col).border = border_all
        if idx % 2 == 1:
            for col in range(1, 16):
                ws2.cell(row=row, column=col).fill = PatternFill("solid", fgColor=LIGHT_GREY)

    n = len(show)
    last_row = n + 1
    cats = Reference(ws2, min_col=1, min_row=2, max_row=last_row)

    def style_line_series(series, color, width_pt):
        series.graphicalProperties.line.solidFill = color
        series.graphicalProperties.line.width = width_pt
        series.marker = Marker(symbol="none")
        series.smooth = False

    def add_pointer_series(chart, col):
        ref = Reference(ws2, min_col=col, min_row=1, max_row=last_row)
        chart.add_data(ref, titles_from_data=True)
        s = chart.series[-1]
        s.marker = Marker(symbol="diamond", size=8)
        s.marker.graphicalProperties.solidFill = RED
        s.marker.graphicalProperties.line.solidFill = RED
        s.graphicalProperties.line.noFill = True
        s.smooth = False
        s.dLbls = DataLabelList()
        s.dLbls.showCatName = True
        s.dLbls.showVal = False
        s.dLbls.showSerName = False
        s.dLbls.showLegendKey = False

    chart1 = LineChart()
    chart1.title = "Overall — Drop % ×100 (blue) · MA10 (orange) · MA50 (green) · pointers = problem episodes"
    chart1.style = 2
    chart1.height = 8.5
    chart1.width = 20
    chart1.x_axis.delete = False
    chart1.y_axis.delete = False
    chart1.displayBlanksAs = "gap"
    for col, color, wpt in [(5, BLUE, 8000), (7, ORANGE, 16000), (9, GREEN, 16000)]:
        ref = Reference(ws2, min_col=col, min_row=1, max_row=last_row)
        chart1.add_data(ref, titles_from_data=True)
        style_line_series(chart1.series[-1], color, wpt)
    chart1.set_categories(cats)
    add_pointer_series(chart1, 13)
    ws.add_chart(chart1, f"A{chart1_row}")

    chart2 = LineChart()
    chart2.title = "Overall — Listening DAU per episode · pointers = problem episodes"
    chart2.style = 2
    chart2.height = 8.5
    chart2.width = 20
    chart2.x_axis.delete = False
    chart2.y_axis.delete = False
    chart2.displayBlanksAs = "gap"
    ldau_ref = Reference(ws2, min_col=3, min_row=1, max_row=last_row)
    chart2.add_data(ldau_ref, titles_from_data=True)
    style_line_series(chart2.series[-1], PURPLE, 16000)
    chart2.set_categories(cats)
    add_pointer_series(chart2, 14)
    ws.add_chart(chart2, f"A{chart2_row}")

    chart3 = LineChart()
    chart3.title = "Retention Index vs The Alpha's Bride (Ep 51 = 100) · pointers = problem episodes"
    chart3.style = 2
    chart3.height = 8.5
    chart3.width = 20
    chart3.x_axis.delete = False
    chart3.y_axis.delete = False
    chart3.displayBlanksAs = "gap"
    show_idx_ref = Reference(ws2, min_col=11, min_row=1, max_row=last_row)
    tab_idx_ref = Reference(ws2, min_col=12, min_row=1, max_row=last_row)
    chart3.add_data(show_idx_ref, titles_from_data=True)
    style_line_series(chart3.series[-1], BLUE, 16000)
    chart3.add_data(tab_idx_ref, titles_from_data=True)
    style_line_series(chart3.series[-1], ORANGE, 16000)
    chart3.set_categories(cats)
    spike_ret_ref = Reference(ws2, min_col=15, min_row=1, max_row=last_row)
    chart3.add_data(spike_ret_ref, titles_from_data=True)
    s = chart3.series[-1]
    s.marker = Marker(symbol="diamond", size=8)
    s.marker.graphicalProperties.solidFill = RED
    s.marker.graphicalProperties.line.solidFill = RED
    s.graphicalProperties.line.noFill = True
    s.smooth = False
    s.dLbls = DataLabelList()
    s.dLbls.showCatName = True
    s.dLbls.showVal = False
    s.dLbls.showSerName = False
    s.dLbls.showLegendKey = False
    ws.add_chart(chart3, f"A{chart3_row}")

    wb.save(out_path)
    return out_path
